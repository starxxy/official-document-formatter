import os
import json
from typing import Dict, Optional
from docx import Document
from docx.shared import Pt, Mm, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from title_detector import TitleDetector


class DocProcessor:
    def __init__(self, config_path: str = None):
        if config_path is None:
            config_path = os.path.join(
                os.path.dirname(os.path.dirname(__file__)),
                "config",
                "official_doc_config.json"
            )
        
        with open(config_path, 'r', encoding='utf-8') as f:
            self.config = json.load(f)
        
        self.title_detector = TitleDetector(config_path)
    
    def load_document(self, file_path: str) -> Document:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文档不存在: {file_path}")
        
        return Document(file_path)
    
    def set_page_format(self, doc: Document) -> None:
        page_config = self.config['page']
        
        section = doc.sections[0]
        
        section.page_width = Mm(page_config['width_mm'])
        section.page_height = Mm(page_config['height_mm'])
        
        section.top_margin = Mm(page_config['top_margin_mm'])
        section.bottom_margin = Mm(page_config['bottom_margin_mm'])
        section.left_margin = Mm(page_config['left_margin_mm'])
        section.right_margin = Mm(page_config['right_margin_mm'])
        
        section.header_distance = Mm(page_config['header_distance_mm'])
        section.footer_distance = Mm(page_config['footer_distance_mm'])
    
    def set_run_font(self, run, font_name: str, font_size: int, bold: bool = False):
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = font_name
        
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    
    def set_paragraph_format(self, paragraph: Paragraph, format_config: Dict):
        pf = paragraph.paragraph_format
        
        if 'first_line_indent_chars' in format_config:
            indent_chars = format_config['first_line_indent_chars']
            if indent_chars > 0:
                pf.first_line_indent = Pt(16 * indent_chars)
            else:
                pf.first_line_indent = Pt(0)
        
        if 'line_spacing_pt' in format_config:
            pf.line_spacing = Pt(format_config['line_spacing_pt'])
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        
        if 'alignment' in format_config:
            alignment_map = {
                'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
                'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
                'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
                'justify': WD_PARAGRAPH_ALIGNMENT.JUSTIFY
            }
            alignment = alignment_map.get(format_config['alignment'], WD_PARAGRAPH_ALIGNMENT.LEFT)
            pf.alignment = alignment
        
        if 'space_before_pt' in format_config:
            pf.space_before = Pt(format_config['space_before_pt'])
        
        if 'space_after_pt' in format_config:
            pf.space_after = Pt(format_config['space_after_pt'])
    
    def apply_font_to_paragraph(self, paragraph: Paragraph, font_config: Dict):
        font_name = font_config['name']
        font_size = font_config['size_pt']
        bold = font_config.get('bold', False)
        
        for run in paragraph.runs:
            self.set_run_font(run, font_name, font_size, bold)
    
    def format_paragraph(self, paragraph: Paragraph, paragraph_type: str = 'body'):
        fonts_config = self.config['fonts']
        paragraph_config = self.config['paragraph']
        
        if paragraph_type == 'main_title':
            font_config = fonts_config['title']
            format_config = paragraph_config['title']
        elif paragraph_type.startswith('heading'):
            level = paragraph_type.replace('heading', '')
            font_config = fonts_config.get(f'heading{level}', fonts_config['body'])
            format_config = paragraph_config.get(f'heading{level}', paragraph_config['body'])
        else:
            font_config = fonts_config['body']
            format_config = paragraph_config['body']
        
        self.apply_font_to_paragraph(paragraph, font_config)
        self.set_paragraph_format(paragraph, format_config)
    
    def process_document(self, doc: Document, main_title_index: Optional[int] = None) -> Dict:
        stats = {
            'total_paragraphs': len(doc.paragraphs),
            'main_titles': 0,
            'headings': {
                'level1': 0,
                'level2': 0,
                'level3': 0,
                'level4': 0
            },
            'body_paragraphs': 0,
            'errors': []
        }
        
        self.set_page_format(doc)
        
        position = 0
        for i, paragraph in enumerate(doc.paragraphs):
            try:
                if not paragraph.text.strip():
                    continue
                
                # 手动指定主标题段落：优先于自动识别
                if main_title_index is not None and i == main_title_index:
                    self.format_paragraph(paragraph, 'main_title')
                    stats['main_titles'] += 1
                    position += 1
                    continue
                
                # 手动指定主标题后，关闭「首段自动判标题」，避免红头被误判
                allow_first = main_title_index is None
                is_main = self.title_detector.is_main_title(paragraph, position, allow_first_paragraph=allow_first)
                
                if is_main:
                    self.format_paragraph(paragraph, 'main_title')
                    stats['main_titles'] += 1
                    position += 1
                    continue
                
                level = self.title_detector.detect_title_level(paragraph)
                
                if level > 0:
                    self.format_paragraph(paragraph, f'heading{level}')
                    level_key = f'level{level}'
                    if level_key in stats['headings']:
                        stats['headings'][level_key] += 1
                else:
                    self.format_paragraph(paragraph, 'body')
                    stats['body_paragraphs'] += 1
                
                position += 1
                    
            except Exception as e:
                error_msg = f"段落 {i} 处理失败: {str(e)}"
                stats['errors'].append(error_msg)
                print(f"警告: {error_msg}")
        
        return stats
    
    def save_document(self, doc: Document, output_path: str) -> bool:
        try:
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            doc.save(output_path)
            return True
        except Exception as e:
            print(f"保存文档失败: {e}")
            return False
    
    def validate_document_format(self, doc: Document) -> Dict:
        validation = {
            'page_format': {},
            'fonts': {},
            'paragraphs': {},
            'issues': []
        }
        
        section = doc.sections[0]
        page_config = self.config['page']
        
        validation['page_format'] = {
            'width': {
                'expected': page_config['width_mm'],
                'actual': section.page_width.mm if section.page_width else None,
                'valid': abs(section.page_width.mm - page_config['width_mm']) < 1 if section.page_width else False
            },
            'height': {
                'expected': page_config['height_mm'],
                'actual': section.page_height.mm if section.page_height else None,
                'valid': abs(section.page_height.mm - page_config['height_mm']) < 1 if section.page_height else False
            },
            'top_margin': {
                'expected': page_config['top_margin_mm'],
                'actual': section.top_margin.mm if section.top_margin else None,
                'valid': abs(section.top_margin.mm - page_config['top_margin_mm']) < 1 if section.top_margin else False
            },
            'left_margin': {
                'expected': page_config['left_margin_mm'],
                'actual': section.left_margin.mm if section.left_margin else None,
                'valid': abs(section.left_margin.mm - page_config['left_margin_mm']) < 1 if section.left_margin else False
            }
        }
        
        title_stats = self.title_detector.get_title_statistics(doc)
        validation['paragraphs'] = {
            'total_paragraphs': len(doc.paragraphs),
            'title_stats': title_stats
        }
        
        hierarchy_validation = self.title_detector.validate_title_hierarchy(doc)
        if not hierarchy_validation['valid']:
            validation['issues'].extend(hierarchy_validation['issues'])
        
        return validation
    
    def create_new_document(self, title: str = None) -> Document:
        doc = Document()
        
        self.set_page_format(doc)
        
        if title:
            title_para = doc.add_paragraph(title)
            self.format_paragraph(title_para, 'main_title')
        
        return doc


def main():
    print("公文格式化 - 文档处理测试")
    print("="*60)
    
    processor = DocProcessor()
    
    print("\n创建新文档测试...")
    doc = processor.create_new_document("测试公文标题")
    
    doc.add_paragraph("这是正文内容，应该使用仿宋_GB2312字体，3号字，首行缩进2字符。")
    
    para1 = doc.add_paragraph("一、一级标题")
    processor.format_paragraph(para1, 'heading1')
    
    para2 = doc.add_paragraph("（一）二级标题")
    processor.format_paragraph(para2, 'heading2')
    
    doc.add_paragraph("这是二级标题下的正文内容。")
    
    output_path = os.path.join(os.path.dirname(__file__), "test_output.docx")
    if processor.save_document(doc, output_path):
        print(f"\n✓ 测试文档已保存: {output_path}")
    else:
        print("\n✗ 测试文档保存失败")
    
    print("\n文档处理模块测试完成！")


if __name__ == "__main__":
    main()
